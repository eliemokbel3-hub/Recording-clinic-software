"""Sample-note learning (note-learning-and-styles plan Phase 3, Task 3.3; D9, D10).

The practitioner chooses one to five of their OWN past notes (``.txt``,
``.docx`` or pasted text); this module reads them INTO MEMORY and derives a
``StyleProfileDraft`` for the Practitioner tab's review screen, which the
tab turns into a ``StyleProfile`` (``build_style_profile``) and writes with
``practitioner_profile.save_style_profile`` only on the practitioner's Save.

What the structure enforces, and only that:

- READ ONLY. ``read_sample_note`` opens a chosen file for reading (a byte
  read for ``.txt``; ``python-docx``'s ``Document`` for ``.docx``, which
  reads the package and never saves) and returns its text; nothing here
  writes, copies, moves or renames a file, and the returned ``SampleNote``
  holds the text and the path, never a copy on disk (D9, C6). Deleting the
  originals is ``delete_sample_files`` — called by the tab ONLY after its
  separate, explicit, default-unticked confirmation that names the paths,
  and it unlinks exactly the listed paths.
- ONE admission control for exemplars (C5). A sentence is kept as an
  exemplar only when THE refusal filter, ``note_config.
  refuse_learning_candidate``, passes every one of its words UNCHANGED —
  original case, original order, the sentence's own first word as the
  ``first_in_segment`` position. A refused sentence is dropped whole, never
  edited to pass (the plan's Excluded item). The learner supplies the
  filter ONE extra (Task 3.7 rule (b)): ``known_common``, the words that
  appear in lowercase anywhere in the chosen notes, so a capitalised
  sentence-opening word the practitioner also writes in lowercase is not
  taken for a name; a shipped abbreviation is never name-like (rule (a),
  inside the filter). Every other capitalised opener, and every
  capitalised mid-sentence word outside the vocabulary, is still refused —
  the safe direction; the review screen shows what was kept and counts
  what was not.
- Shorthand by controlled vocabulary (D10). A token that is EXACTLY in the
  shipped ``note_config.CLINICAL_ABBREVIATIONS`` (case-preserving — ``as``
  is a word, ``AS`` the abbreviation) is ``recognised`` and saved; an
  abbreviation-SHAPED token that is not in it (all capitals, a ``Cx``-like
  form, a ``C/O``-like form — never a token with a digit) is listed as
  ``unrecognised`` and saved only when the practitioner ticks it;
  ``build_style_profile`` refuses a shorthand token or an exemplar the draft
  did not offer, so the review can only remove, never add or edit.
- Nothing here logs. ``SampleNote.sample_text``, ``recognised_shorthand``
  and ``unrecognised_shorthand`` are log-tripwire signatures
  (``logging_setup``) beside ``exemplar_text`` (C9).

Heuristics, NOT controls (a wrong guess costs a review glance, never a
patient detail): a heading is a short line (or a line's prefix before a
colon) whose normalised text matches one of the canonical section titles or
the alias table below; the section order is the order of first appearance
across the notes in the order given; a sentence is a run of at least
``_MIN_SENTENCE_WORDS`` words split on terminal punctuation or a line
break; the style measures are simple counts.
"""

from __future__ import annotations

import io
import re
import zipfile
import zlib
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Final, NamedTuple

from pydantic import ValidationError

from scribe_desktop.note import CANONICAL_SECTIONS, NoteSectionKey, strip_token_punctuation
from scribe_desktop.note_config import (
    # ``_no_control_chars`` is package-private by name, shared deliberately
    # (the ``ui.models`` convention): ONE config-text validator, applied to a
    # heading label at parse time so the draft never carries a label the
    # ``StyleProfile`` model would refuse at Save.
    CLINICAL_ABBREVIATIONS,
    MAX_CONFIG_LABEL_CHARS,
    MAX_SAMPLE_NOTES,
    MAX_STYLE_EXEMPLARS,
    RefusalClass,
    StyleExemplar,
    StyleMeasures,
    StylePerson,
    StyleProfile,
    StyleTense,
    _no_control_chars,
    refuse_learning_candidate,
)
from scribe_desktop.practitioner_profile import ConsentRecord

# A sample note is a clinical note, not a document archive: these bounds
# keep one read in memory small and refuse a mis-chosen file loudly.
MAX_SAMPLE_NOTE_BYTES: Final = 2 * 1024 * 1024
MAX_SAMPLE_NOTE_CHARS: Final = 200_000
SAMPLE_NOTE_SUFFIXES: Final[tuple[str, ...]] = (".txt", ".docx")
# Peer rounds 17 / 18 (PR-MED-020, PR-LOW-027): a ``.docx`` is a zip package
# that python-docx inflates member by member IN FULL before any text is seen,
# so the compressed-size bound above is not a bound on memory — and neither
# are the sizes the central directory DECLARES on their own: an unbounded
# ``ZipFile.read()`` asks the decompressor for up to ~2 GiB of output per
# call and only THEN truncates what it returns to the declared size, so a
# lying directory still buys the allocation. The guarantee this module gives
# is therefore its OWN, in two parts. (1) ALLOCATION: it inflates every member
# itself through ``ZipFile.open(...).read(cap + 1)`` — never an unbounded read
# (pinned by a test that records every ``n`` this module asks for) — and the
# stdlib passes ``max(n, 4096)`` (its ``MIN_READ_SIZE`` floor) to the
# decompressor as ``max_length``, so the output allocated per ``decompress``
# call is at most ``MAX_DOCX_MEMBER_BYTES + 1`` bytes (the sentinel byte
# included; the 4 KiB floor matters only below it), whatever the directory
# declared. (2) BYTES: the stdlib truncates what it RETURNS to the declared
# size and CRC-checks the member at its end, so an understated size yields at
# most the declared bytes (already ≤ the cap after the directory check) and,
# unless the CRC was forged to match, a named ``BadZipFile`` refusal — the two
# over-cap refusals in the read loop are DEFENSIVE, reachable only if the
# stdlib ever returned more than it declared. The directory is checked first
# (member count, declared per-member size and total, STORED / DEFLATED only —
# Word writes nothing else) as a fast refusal that inflates nothing; the
# bounded bytes are then re-zipped STORED into memory and THAT package — never
# the chosen file — is what ``Document`` parses. Residue, named: the
# compressed input read per call is bounded by ``MAX_SAMPLE_NOTE_BYTES``; the
# decompressor's own window and the ≤ ``MAX_DOCX_DECLARED_BYTES`` re-zip
# buffer are held while the document is parsed; python-docx's XML parse of a
# member is bounded only by that member's cap. A note with a few embedded
# images stays far inside these bounds; a bomb or a mis-chosen archive is
# refused by name.
MAX_DOCX_MEMBERS: Final = 512
MAX_DOCX_MEMBER_BYTES: Final = 8 * 1024 * 1024
MAX_DOCX_DECLARED_BYTES: Final = 16 * 1024 * 1024
_DOCX_METHODS: Final[frozenset[int]] = frozenset({zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED})
# The review screen lists at most this many unrecognised tokens.
MAX_UNRECOGNISED_SHORTHAND: Final = 100

_MIN_SENTENCE_WORDS: Final = 3
_MAX_HEADING_WORDS: Final = 6
_SENTENCE_SPLIT_RE: Final = re.compile(r"(?<=[.!?;])\s+")
_ALL_CAPS_RE: Final = re.compile(r"^[A-Z]{2,8}$")
_X_FORM_RE: Final = re.compile(r"^[A-Z]{1,4}x$")
_SLASH_FORM_RE: Final = re.compile(r"^[A-Z]{1,3}/[A-Z]{1,3}$")
_HEADING_STRIP_RE: Final = re.compile(r"[^a-z0-9/& ]+")

# Heading aliases (lower-case, punctuation-stripped) beside the canonical
# titles. A heuristic table: an unlisted heading simply leaves the text
# unsectioned (and its sentences are not offered as exemplars).
_HEADING_ALIASES: Final[Mapping[str, NoteSectionKey]] = {
    "c/o": "presenting_complaint",
    "co": "presenting_complaint",
    "complaint": "presenting_complaint",
    "presenting complaint": "presenting_complaint",
    "pc": "presenting_complaint",
    "subjective": "presenting_complaint",
    "s": "presenting_complaint",
    "hpc": "history_presenting_complaint",
    "history": "history_presenting_complaint",
    "hx": "history_presenting_complaint",
    "history of presenting complaint": "history_presenting_complaint",
    "progress": "progress_since_last_visit",
    "progress since last visit": "progress_since_last_visit",
    "since last visit": "progress_since_last_visit",
    "pmh": "past_medical_history",
    "pmhx": "past_medical_history",
    "past history": "past_medical_history",
    "past medical history": "past_medical_history",
    "medical history": "past_medical_history",
    "red flags": "red_flags_screening",
    "red flags screening": "red_flags_screening",
    "o/e": "objective_examination",
    "oe": "objective_examination",
    "objective": "objective_examination",
    "objective examination": "objective_examination",
    "examination": "objective_examination",
    "exam": "objective_examination",
    "o": "objective_examination",
    "outcome measures": "outcome_measures",
    "outcomes": "outcome_measures",
    "assessment": "assessment",
    "impression": "assessment",
    "a": "assessment",
    "diagnosis": "diagnosis",
    "dx": "diagnosis",
    "treatment": "treatment_performed",
    "treatment performed": "treatment_performed",
    "tx": "treatment_performed",
    "rx": "treatment_performed",
    "response": "response_to_treatment",
    "response to treatment": "response_to_treatment",
    "post treatment": "response_to_treatment",
    "post tx": "response_to_treatment",
    "advice": "advice_home_exercise",
    "advice and home exercise": "advice_home_exercise",
    "home exercise": "advice_home_exercise",
    "hep": "advice_home_exercise",
    "exercises": "advice_home_exercise",
    "plan": "management_plan",
    "management": "management_plan",
    "management plan": "management_plan",
    "p": "management_plan",
    "consent": "consent",
    "referral": "referrals_investigations",
    "referrals": "referrals_investigations",
    "investigations": "referrals_investigations",
    "referrals and investigations": "referrals_investigations",
    "precautions": "precautions_contraindications",
    "contraindications": "precautions_contraindications",
    "precautions and contraindications": "precautions_contraindications",
    "follow up": "follow_up_review",
    "followup": "follow_up_review",
    "f/u": "follow_up_review",
    "review": "follow_up_review",
    "follow-up and review": "follow_up_review",
    "follow up and review": "follow_up_review",
    "r/v": "follow_up_review",
}

_FIRST_PERSON: Final[frozenset[str]] = frozenset(
    "i i'm i've i'll my me we we're we've our us".split()
)
_THIRD_PERSON: Final[frozenset[str]] = frozenset(
    "patient pt client he she they his her their him them".split()
)
_PAST_MARKERS: Final[frozenset[str]] = frozenset(
    "was were had did reported advised performed applied presented noted".split()
)
_PRESENT_MARKERS: Final[frozenset[str]] = frozenset(
    "is are has reports presents complains denies feels continues advise apply".split()
)


class SampleNoteError(Exception):
    """A sample note could not be read or the draft could not be built."""


@dataclass(frozen=True)
class SampleNote:
    """One note read into memory: its text (never rendered by ``repr``) and
    the path it came from — None for pasted text, which has no file to
    offer for deletion."""

    sample_text: str = field(repr=False)
    source_path: Path | None = None


class _Heading(NamedTuple):
    key: NoteSectionKey
    label: str


def _vocabulary() -> frozenset[str]:
    """The shipped vocabulary as written: a token is recognised by EXACT,
    case-preserving match only (Task 3.7 rule (a) is the same match) —
    ``as`` is a common word, ``AS`` the abbreviation, and only the latter
    counts."""
    return frozenset(CLINICAL_ABBREVIATIONS)


def _canonical_titles() -> Mapping[str, NoteSectionKey]:
    titles: dict[str, NoteSectionKey] = {}
    for section in CANONICAL_SECTIONS:
        titles[_normalise_heading(section.title)] = section.key
        titles[section.key.replace("_", " ")] = section.key
    return titles


def _normalise_heading(text: str) -> str:
    lowered = _HEADING_STRIP_RE.sub(" ", text.casefold())
    return " ".join(lowered.split())


_TITLE_KEYS: Final[Mapping[str, NoteSectionKey]] = _canonical_titles()


def _match_heading(line: str) -> tuple[_Heading, str] | None:
    """A heading line (``Assessment``), or a heading prefix with its body
    (``Plan: continue HEP``) → the section and the remainder of the line."""
    head, colon, rest = line.partition(":")
    head = head.strip()
    if not head or len(head.split()) > _MAX_HEADING_WORDS:
        return None
    if not colon and len(line.split()) > _MAX_HEADING_WORDS:
        return None
    if len(head) > MAX_CONFIG_LABEL_CHARS:
        return None
    try:
        _no_control_chars(head)  # the label must be a legal StyleProfile heading
    except ValueError:
        return None  # fail toward "not a heading": its lines are unsectioned, never kept
    normalised = _normalise_heading(head)
    key = _HEADING_ALIASES.get(normalised) or _TITLE_KEYS.get(normalised)
    if key is None:
        return None
    return _Heading(key, head), rest.strip()


# --- reading -----------------------------------------------------------------


def _decode(blob: bytes, name: str) -> str:
    try:
        return blob.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        return blob.decode("cp1252")
    except UnicodeDecodeError:
        raise SampleNoteError(f"could not read {name}: not a text file") from None


# The standard library's raise set at ``ZipFile(path)`` + ``infolist()`` +
# ``open()`` + the member read (peer rounds 18 / 19, PR-LOW-028 / PR-LOW-030),
# named exactly so a bug in this module still surfaces as a bug:
# ``BadZipFile`` (a corrupt directory or a CRC mismatch; a ``struct.error``
# from a corrupt extra field is re-raised by the stdlib as this), ``OSError``
# (the read itself), ``UnicodeDecodeError`` (CPython's ``_RealGetContents``
# decodes a member name with ``filename.decode('utf-8')`` under
# general-purpose flag 0x800) and ``NotImplementedError`` (the same function
# on ``extract_version > MAX_EXTRACT_VERSION``; also ``open()`` on a
# compression method the stdlib lacks — refused earlier here by method).
# Two more input failures are refused elsewhere: an ENCRYPTED member (flag
# bit 0x1, which ``open()`` would meet as ``RuntimeError`` "password
# required") is refused by name at the directory stage before any ``open()``,
# and corrupt DEFLATE data (``zlib.error`` from the decompressor inside the
# member read) is translated by the two-statement ``try`` around that read.
# Residue: an un-enumerated stdlib raise on a future Python reaches the tab
# as a bug, which is the intended direction.
_ZIP_METADATA_ERRORS: Final = (
    zipfile.BadZipFile,
    OSError,
    UnicodeDecodeError,
    NotImplementedError,
)
_ZIP_ENCRYPTED_FLAG: Final = 0x1


def _too_large(path: Path, what: str, size: int, limit: int) -> SampleNoteError:
    return SampleNoteError(
        f"{path.name} is too large to be a note: {what} is more than {size:,} bytes when "
        f"expanded (the limit is {limit:,})"
    )


def _bounded_docx_package(path: Path) -> io.BytesIO:
    """PR-MED-020 / PR-LOW-027: the chosen file's members inflated by THIS
    module under hard caps and re-zipped STORED into memory — the only
    package python-docx ever parses. The central directory is read first
    (member count, each member's DECLARED size, their total, the method — a
    fast refusal that inflates nothing; an encrypted member refused here by
    its flag bit); then every member is read through ``ZipFile.open`` with
    ``read(cap + 1)`` — the request the stdlib rounds to its 4 KiB floor and
    passes as ``max_length``, so the decompressor's allowance per call is
    ``max(cap + 1, 4096)`` bytes whatever the directory declared (the stdlib
    truncates the returned bytes to the declared size and refuses a CRC
    mismatch, so the over-cap branches below are defensive; corrupt DEFLATE
    data is a named refusal). A file that is not a zip, or whose directory
    the stdlib cannot parse, is refused by name too."""
    total = 0
    bounded = io.BytesIO()
    try:
        with (
            zipfile.ZipFile(path) as archive,
            zipfile.ZipFile(bounded, "w", zipfile.ZIP_STORED) as copy,
        ):
            infos = archive.infolist()
            if len(infos) > MAX_DOCX_MEMBERS:
                raise SampleNoteError(
                    f"{path.name} is not a note: the document package holds "
                    f"{len(infos):,} parts (the limit is {MAX_DOCX_MEMBERS})"
                )
            for info in infos:
                if info.compress_type not in _DOCX_METHODS:
                    raise SampleNoteError(
                        f"could not read {path.name}: not a Word document (a part uses "
                        f"compression method {info.compress_type})"
                    )
                if info.flag_bits & _ZIP_ENCRYPTED_FLAG:
                    # PR-LOW-030: Word never encrypts an OPC part; ``open()``
                    # would raise ``RuntimeError`` for the missing password.
                    raise SampleNoteError(
                        f"could not read {path.name}: not a Word document (a part is "
                        "encrypted)"
                    )
                if info.file_size > MAX_DOCX_MEMBER_BYTES:
                    raise _too_large(
                        path, "one part of the document package", info.file_size,
                        MAX_DOCX_MEMBER_BYTES,
                    )
                total += info.file_size
            if total > MAX_DOCX_DECLARED_BYTES:
                raise _too_large(
                    path, "the document package", total, MAX_DOCX_DECLARED_BYTES
                )
            total = 0
            for info in infos:
                if info.is_dir():
                    continue
                try:
                    with archive.open(info) as member:
                        # ``n`` is the request the stdlib rounds to its 4 KiB
                        # floor and passes as ``max_length`` (never ``read()``
                        # / ``read(-1)`` here); see the constants' comment.
                        data = member.read(MAX_DOCX_MEMBER_BYTES + 1)
                except zlib.error as exc:  # PR-LOW-030: corrupt DEFLATE data
                    raise SampleNoteError(
                        f"could not read {path.name}: not a Word document "
                        f"({type(exc).__name__})"
                    ) from None
                if len(data) > MAX_DOCX_MEMBER_BYTES:  # defensive, see the constants
                    raise _too_large(
                        path, "one part of the document package", MAX_DOCX_MEMBER_BYTES,
                        MAX_DOCX_MEMBER_BYTES,
                    )
                total += len(data)
                if total > MAX_DOCX_DECLARED_BYTES:
                    raise _too_large(
                        path, "the document package", MAX_DOCX_DECLARED_BYTES,
                        MAX_DOCX_DECLARED_BYTES,
                    )
                copy.writestr(info.filename, data)
                del data
    except _ZIP_METADATA_ERRORS as exc:
        raise SampleNoteError(
            f"could not read {path.name}: not a Word document ({type(exc).__name__})"
        ) from None
    bounded.seek(0)
    return bounded


def _read_docx(path: Path) -> str:
    package = _bounded_docx_package(path)  # before the library is even imported
    try:
        import docx
    except ImportError:
        raise SampleNoteError(
            "reading a .docx note needs python-docx, which is not installed - paste the "
            "text instead"
        ) from None
    try:
        document = docx.Document(package)
    except Exception as exc:  # noqa: BLE001 - python-docx raises library-specific types
        raise SampleNoteError(
            f"could not read {path.name}: not a Word document ({type(exc).__name__})"
        ) from None
    finally:
        package.close()
    # The text is bounded WHILE it is collected (PR-MED-020), not only after:
    # a package inside the declared-size bound can still carry more text
    # than a note.
    lines: list[str] = []
    collected = 0

    def take(text: str) -> None:
        nonlocal collected
        collected += len(text)
        if collected > MAX_SAMPLE_NOTE_CHARS:
            raise SampleNoteError(
                f"{path.name} is too long to be a note (more than "
                f"{MAX_SAMPLE_NOTE_CHARS:,} characters of text)"
            )
        lines.append(text)

    for paragraph in document.paragraphs:
        take(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    take(paragraph.text)
    return "\n".join(lines)


def _bounded(text: str, name: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if len(text) > MAX_SAMPLE_NOTE_CHARS:
        raise SampleNoteError(
            f"{name} is too long to be a note ({len(text):,} characters; the limit is "
            f"{MAX_SAMPLE_NOTE_CHARS:,})"
        )
    if not text.strip():
        raise SampleNoteError(f"{name} holds no text to learn from")
    return text


def read_sample_note(source: Path | str) -> SampleNote:
    """A note into memory. A ``str`` is pasted text; a ``Path`` is a ``.txt``
    or ``.docx`` file, read once and never written, copied or moved.
    ``SampleNoteError`` names the reason otherwise (unknown suffix, too
    large, unreadable, empty)."""
    if isinstance(source, str):
        return SampleNote(_bounded(source, "the pasted text"), None)
    suffix = source.suffix.casefold()
    if suffix not in SAMPLE_NOTE_SUFFIXES:
        raise SampleNoteError(
            f"{source.name} is not a .txt or .docx file (PDF notes are not read - paste "
            "the text instead)"
        )
    try:
        size = source.stat().st_size
    except OSError as exc:
        raise SampleNoteError(f"could not read {source.name}: {exc}") from None
    if size > MAX_SAMPLE_NOTE_BYTES:
        raise SampleNoteError(
            f"{source.name} is too large to be a note ({size:,} bytes; the limit is "
            f"{MAX_SAMPLE_NOTE_BYTES:,})"
        )
    if suffix == ".docx":
        text = _read_docx(source)
    else:
        try:
            blob = source.read_bytes()
        except OSError as exc:
            raise SampleNoteError(f"could not read {source.name}: {exc}") from None
        text = _decode(blob, source.name)
    return SampleNote(_bounded(text, source.name), source)


# --- learning ----------------------------------------------------------------


@dataclass(frozen=True)
class StyleProfileDraft:
    """What the learner derived, for the review screen: everything a
    ``StyleProfile`` will hold except the consent record and the time, plus
    the split shorthand (D10), the source paths (for the delete-originals
    step) and the counts of what was NOT kept, so the screen can say so."""

    section_order: tuple[NoteSectionKey, ...]
    heading_labels: Mapping[NoteSectionKey, str]
    recognised_shorthand: tuple[str, ...]
    unrecognised_shorthand: tuple[str, ...]
    measures: StyleMeasures
    exemplars: tuple[StyleExemplar, ...] = field(repr=False)
    source_count: int
    source_paths: tuple[Path, ...]
    refused_exemplars: Mapping[RefusalClass, int]
    dropped_exemplars: int


class _Parsed(NamedTuple):
    headings: tuple[_Heading, ...]
    sections: tuple[tuple[NoteSectionKey | None, str], ...]  # (section, one line)


def _parse(text: str) -> _Parsed:
    headings: list[_Heading] = []
    seen: set[NoteSectionKey] = set()
    current: NoteSectionKey | None = None
    lines: list[tuple[NoteSectionKey | None, str]] = []
    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            continue
        matched = _match_heading(line)
        if matched is not None:
            heading, rest = matched
            current = heading.key
            if heading.key not in seen:
                seen.add(heading.key)
                headings.append(heading)
            if rest:
                lines.append((current, rest))
            continue
        lines.append((current, line))
    return _Parsed(tuple(headings), tuple(lines))


def _sentences(line: str) -> list[str]:
    return [part.strip() for part in _SENTENCE_SPLIT_RE.split(line) if part.strip()]


def _words(text: str) -> list[str]:
    """The line's words with the ONE punctuation rule applied and the case
    kept (``note.strip_token_punctuation``)."""
    stripped = (strip_token_punctuation(raw) for raw in text.split())
    return [word for word in stripped if word]


def _evidence(notes: Sequence[SampleNote]) -> frozenset[str]:
    """Task 3.7 rule (b)'s evidence: every word that appears in LOWERCASE
    anywhere in the notes (punctuation-stripped). Passed to the refusal
    filter as ``known_common`` so a capitalised sentence-opening word the
    practitioner also writes in lowercase ("Review …" / "… for review") is
    not taken for a name; a word seen only capitalised is never admitted."""
    seen: set[str] = set()
    for note in notes:
        for raw in note.sample_text.split():
            word = strip_token_punctuation(raw)
            if word and word == word.lower() and any(ch.isalpha() for ch in word):
                seen.add(word)
    return frozenset(seen)


def _shorthand_shape(token: str) -> bool:
    if any(ch.isdigit() for ch in token):
        return False
    return bool(
        _ALL_CAPS_RE.match(token) or _X_FORM_RE.match(token) or _SLASH_FORM_RE.match(token)
    )


def _person(person_first: int, person_third: int) -> StylePerson:
    if person_first == 0 and person_third == 0:
        return "unknown"
    if person_first and person_third:
        low, high = sorted((person_first, person_third))
        if low / high >= 0.25:
            return "mixed"
    return "first" if person_first > person_third else "third"


def _tense(past: int, present: int) -> StyleTense:
    if past == 0 and present == 0:
        return "unknown"
    if past and present:
        low, high = sorted((past, present))
        if low / high >= 0.25:
            return "mixed"
    return "past" if past > present else "present"


def _refuse(sentence: str, evidence: frozenset[str]) -> RefusalClass | None:
    """THE refusal filter over the sentence's raw words, unchanged: its
    first word is the sentence's real first word, and the notes' lowercase
    evidence is the only extra the learner supplies (Task 3.7)."""
    return refuse_learning_candidate(
        sentence.split(), first_in_segment=True, following=(), known_common=evidence
    )


def _spread(by_section: Mapping[NoteSectionKey, list[StyleExemplar]]) -> list[StyleExemplar]:
    """At most ``MAX_STYLE_EXEMPLARS``, taken round-robin across the sections
    in the order given so one long section cannot crowd the others out."""
    chosen: list[StyleExemplar] = []
    queues = [list(items) for items in by_section.values()]
    while len(chosen) < MAX_STYLE_EXEMPLARS and any(queues):
        for queue in queues:
            if queue and len(chosen) < MAX_STYLE_EXEMPLARS:
                chosen.append(queue.pop(0))
    return chosen


def learn_style_profile(notes: Sequence[SampleNote]) -> StyleProfileDraft:
    """Derive the draft from 1–``MAX_SAMPLE_NOTES`` notes (``SampleNoteError``
    outside that range). Pure over the texts: reads no file, writes
    nothing."""
    if not 1 <= len(notes) <= MAX_SAMPLE_NOTES:
        raise SampleNoteError(
            f"learn from 1 to {MAX_SAMPLE_NOTES} notes at a time ({len(notes)} given)"
        )
    vocabulary = _vocabulary()
    evidence = _evidence(notes)
    order: list[NoteSectionKey] = []
    labels: dict[NoteSectionKey, str] = {}
    recognised: dict[str, None] = {}
    unrecognised: dict[str, None] = {}
    candidates: dict[NoteSectionKey, list[StyleExemplar]] = {}
    refused: dict[RefusalClass, int] = {}
    dropped = 0
    word_total = 0
    abbreviation_total = 0
    sentence_words: list[int] = []
    person_first = person_third = past = present = 0

    for note in notes:
        parsed = _parse(note.sample_text)
        for heading in parsed.headings:
            if heading.key not in labels:
                order.append(heading.key)
                labels[heading.key] = heading.label
        for section, line in parsed.sections:
            words = _words(line)
            word_total += len(words)
            for word in words:
                lowered = word.casefold()
                if lowered in _FIRST_PERSON:
                    person_first += 1
                elif lowered in _THIRD_PERSON:
                    person_third += 1
                if lowered in _PAST_MARKERS or (len(lowered) > 3 and lowered.endswith("ed")):
                    past += 1
                elif lowered in _PRESENT_MARKERS:
                    present += 1
                if word in vocabulary:
                    abbreviation_total += 1
                    recognised.setdefault(word, None)
                elif _shorthand_shape(word):
                    abbreviation_total += 1
                    unrecognised.setdefault(word, None)
            for sentence in _sentences(line):
                count = len(sentence.split())
                if count < _MIN_SENTENCE_WORDS:
                    continue
                sentence_words.append(count)
                if section is None:
                    dropped += 1
                    continue
                refusal = _refuse(sentence, evidence)
                if refusal is not None:
                    refused[refusal] = refused.get(refusal, 0) + 1
                    continue
                try:
                    exemplar = StyleExemplar(section_key=section, exemplar_text=sentence)
                except ValidationError:
                    dropped += 1  # over the exemplar length, or a control character
                    continue
                candidates.setdefault(section, []).append(exemplar)

    kept = _spread(candidates)
    dropped += sum(len(items) for items in candidates.values()) - len(kept)
    mean_words = sum(sentence_words) / len(sentence_words) if sentence_words else 0.0
    ratio = abbreviation_total / word_total if word_total else 0.0
    measures = StyleMeasures(
        mean_sentence_words=round(mean_words, 2),
        abbreviation_ratio=round(min(ratio, 1.0), 4),
        person=_person(person_first, person_third),
        tense=_tense(past, present),
    )
    return StyleProfileDraft(
        section_order=tuple(order),
        heading_labels=dict(labels),
        recognised_shorthand=tuple(recognised),
        unrecognised_shorthand=tuple(unrecognised)[:MAX_UNRECOGNISED_SHORTHAND],
        measures=measures,
        exemplars=tuple(kept),
        source_count=len(notes),
        source_paths=tuple(note.source_path for note in notes if note.source_path is not None),
        refused_exemplars=dict(refused),
        dropped_exemplars=dropped,
    )


def build_style_profile(
    draft: StyleProfileDraft,
    *,
    kept_unrecognised: Sequence[str],
    kept_exemplars: Sequence[StyleExemplar],
    consent: ConsentRecord,
    learned_at: datetime,
) -> StyleProfile:
    """The profile the tab saves after review: every recognised token, the
    ticked unrecognised tokens, and the exemplars the practitioner left in
    place. The review may only REMOVE: a token the draft did not list as
    unrecognised, or an exemplar that is not one of the draft's own (text and
    section equal), is refused with ``SampleNoteError`` — the words that
    passed the filter are the words that are saved (C5, C6)."""
    offered = set(draft.unrecognised_shorthand)
    for token in kept_unrecognised:
        if token not in offered:
            raise SampleNoteError("a shorthand token that was not offered cannot be kept")
    offered_exemplars = set(draft.exemplars)
    for exemplar in kept_exemplars:
        if exemplar not in offered_exemplars:
            raise SampleNoteError("an example sentence that was not offered cannot be kept")
    shorthand: dict[str, None] = dict.fromkeys(draft.recognised_shorthand)
    for token in kept_unrecognised:
        shorthand.setdefault(token, None)
    seen: set[StyleExemplar] = set()
    exemplars: list[StyleExemplar] = []
    for exemplar in kept_exemplars:
        if exemplar not in seen:
            seen.add(exemplar)
            exemplars.append(exemplar)
    try:
        return StyleProfile(
            learned_at=learned_at,
            consent=consent,
            section_order=draft.section_order,
            heading_labels=dict(draft.heading_labels),
            shorthand=tuple(shorthand),
            measures=draft.measures,
            exemplars=tuple(exemplars),
            source_count=draft.source_count,
        )
    except ValidationError as exc:
        # ``hide_input_in_errors``: the message names fields, never text.
        raise SampleNoteError(f"the learned style could not be built: {exc}") from None


# --- deleting the originals (D9: a separate explicit confirmation) --------------


class DeletionReport(NamedTuple):
    deleted: tuple[Path, ...]
    failed: tuple[tuple[Path, str], ...]


def delete_sample_files(paths: Sequence[Path]) -> DeletionReport:
    """Unlink EXACTLY the listed files — nothing beside them, nothing under
    them (a directory is refused, not walked). Called by the tab only after
    the practitioner's separate confirmation naming these paths; a file
    already gone counts as deleted; any other failure is reported per path
    and the rest are still attempted."""
    deleted: list[Path] = []
    failed: list[tuple[Path, str]] = []
    for path in paths:
        if path.is_dir():
            failed.append((path, "is a folder, not a note"))
            continue
        try:
            path.unlink(missing_ok=True)
        except OSError as exc:
            failed.append((path, str(exc)))
            continue
        deleted.append(path)
    return DeletionReport(tuple(deleted), tuple(failed))


__all__ = [
    "MAX_DOCX_DECLARED_BYTES",
    "MAX_DOCX_MEMBERS",
    "MAX_DOCX_MEMBER_BYTES",
    "MAX_SAMPLE_NOTE_BYTES",
    "MAX_SAMPLE_NOTE_CHARS",
    "MAX_UNRECOGNISED_SHORTHAND",
    "SAMPLE_NOTE_SUFFIXES",
    "DeletionReport",
    "SampleNote",
    "SampleNoteError",
    "StyleProfileDraft",
    "build_style_profile",
    "delete_sample_files",
    "learn_style_profile",
    "read_sample_note",
]
